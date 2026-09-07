import os
import io
import datetime
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

from database import get_all_employees_no_blob, get_attendance_logs

# Load environment variables
load_dotenv()

def get_email_config():
    sender = os.getenv('REPORT_SENDER_EMAIL', 'keyaattendence@gmail.com')
    password = os.getenv('REPORT_SENDER_PASSWORD', 'fymgxevokazgkzst')
    receiver = os.getenv('REPORT_RECEIVER_EMAIL', 'keyafusion@gmail.com')
    raw_schedule = os.getenv('REPORT_SCHEDULE_TIMES') or os.getenv('REPORT_SCHEDULE_TIME') or '09:05,20:00'
    schedule_times = [t.strip() for t in raw_schedule.split(',') if t.strip()]
    
    formatted_labels = []
    for t in schedule_times:
        try:
            t_obj = datetime.datetime.strptime(t, "%H:%M")
            formatted_labels.append(t_obj.strftime("%I:%M %p"))
        except Exception:
            formatted_labels.append(t)

    return {
        'sender': sender,
        'password': password,
        'receiver': receiver,
        'schedule_times': schedule_times,
        'schedule_time_display': " & ".join(formatted_labels)
    }

def get_daily_attendance_data(target_date=None):
    """
    Collects attendance records and statistics for the given target date (YYYY-MM-DD).
    Defaults to today's date.
    """
    if not target_date:
        target_date = datetime.date.today().strftime('%Y-%m-%d')

    employees = get_all_employees_no_blob(active_only=True)
    logs = get_attendance_logs(date=target_date)

    # Index logs by employee_id
    log_map = {}
    for log in logs:
        eid = log['employee_id']
        log_map[eid] = log

    records = []
    present_count = 0
    absent_count = 0
    leave_count = 0

    for emp in employees:
        eid = emp['employee_id']
        name = emp.get('name', 'N/A')
        dept = emp.get('department', '-')
        phone = emp.get('phone', '-')

        log = log_map.get(eid)
        if log:
            login_t = log.get('login_time', '')
            logout_t = log.get('logout_time', '') or '-'

            if login_t == 'Absent':
                status = 'Absent'
                absent_count += 1
            elif login_t in ['Sick Leave', 'Paid Leave', 'Company Holiday']:
                status = login_t
                leave_count += 1
            elif login_t:
                status = 'Present'
                present_count += 1
            else:
                status = 'Absent'
                absent_count += 1
        else:
            login_t = '-'
            logout_t = '-'
            status = 'Absent'
            absent_count += 1

        records.append({
            'employee_id': eid,
            'name': name,
            'department': dept,
            'phone': phone,
            'date': target_date,
            'login_time': login_t,
            'logout_time': logout_t,
            'status': status
        })

    total_employees = len(employees)
    attendance_rate = round((present_count / total_employees * 100), 1) if total_employees > 0 else 0.0

    summary = {
        'date': target_date,
        'total_employees': total_employees,
        'present_count': present_count,
        'absent_count': absent_count,
        'leave_count': leave_count,
        'attendance_rate': attendance_rate
    }

    return summary, records

def generate_daily_excel_bytes(summary, records):
    """
    Generates a beautifully styled Excel workbook for the daily attendance as bytes.
    """
    data = []
    for r in records:
        data.append({
            'Employee ID': r['employee_id'],
            'Employee Name': r['name'],
            'Department': r['department'],
            'Phone': r['phone'],
            'Date': r['date'],
            'Login Time': r['login_time'],
            'Logout Time': r['logout_time'],
            'Status': r['status']
        })

    df = pd.DataFrame(data)
    if df.empty:
        df = pd.DataFrame(columns=['Employee ID', 'Employee Name', 'Department', 'Phone', 'Date', 'Login Time', 'Logout Time', 'Status'])

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Daily Attendance', index=False)
        worksheet = writer.sheets['Daily Attendance']

        # Styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        center_align = Alignment(horizontal='center', vertical='center')

        present_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        leave_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        paid_fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid")
        holiday_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Header formatting
        for col_num, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            col_letter = get_column_letter(col_num)
            worksheet.column_dimensions[col_letter].width = 20

        worksheet.row_dimensions[1].height = 28

        # Row formatting
        for r_idx in range(2, len(df) + 2):
            worksheet.row_dimensions[r_idx].height = 24
            status_val = str(worksheet.cell(row=r_idx, column=8).value or '')

            for c_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=r_idx, column=c_idx)
                cell.alignment = center_align
                cell.border = thin_border

                if status_val == 'Present':
                    if c_idx == 8:
                        cell.fill = present_fill
                        cell.font = Font(color="065F46", bold=True)
                elif status_val == 'Absent':
                    if c_idx == 8:
                        cell.fill = absent_fill
                        cell.font = Font(color="991B1B", bold=True)
                elif status_val == 'Sick Leave':
                    if c_idx == 8:
                        cell.fill = leave_fill
                        cell.font = Font(color="92400E", bold=True)
                elif status_val == 'Paid Leave':
                    if c_idx == 8:
                        cell.fill = paid_fill
                        cell.font = Font(color="6B21A8", bold=True)
                elif 'Holiday' in status_val:
                    if c_idx == 8:
                        cell.fill = holiday_fill
                        cell.font = Font(color="3730A3", bold=True)

    output.seek(0)
    return output.read()

def generate_html_email_body(summary, records, session_name=None):
    """
    Generates a high-quality responsive HTML email template with branding, metric cards, and table.
    """
    date_obj = datetime.datetime.strptime(summary['date'], '%Y-%m-%d')
    formatted_date = date_obj.strftime('%A, %d %B %Y')

    if not session_name:
        current_hour = datetime.datetime.now().hour
        session_name = "Evening Summary (8:00 PM)" if current_hour >= 16 else "Morning Check-In (9:05 AM)"

    rows_html = ""
    for r in records:
        status = r['status']
        if status == 'Present':
            badge_style = "background-color:#d1fae5; color:#065f46; border:1px solid #a7f3d0;"
        elif status == 'Absent':
            badge_style = "background-color:#fee2e2; color:#991b1b; border:1px solid #fecaca;"
        elif status == 'Sick Leave':
            badge_style = "background-color:#fef3c7; color:#92400e; border:1px solid #fde68a;"
        elif status == 'Paid Leave':
            badge_style = "background-color:#f3e8ff; color:#6b21a8; border:1px solid #e9d5ff;"
        else:
            badge_style = "background-color:#e0e7ff; color:#3730a3; border:1px solid #c7d2fe;"

        login_display = r['login_time'] if r['login_time'] != '-' and r['login_time'] != 'Absent' else '-'
        logout_display = r['logout_time'] if r['logout_time'] != '-' and r['logout_time'] != 'Absent' else '-'

        rows_html += f"""
        <tr style="border-bottom:1px solid #f1f5f9; text-align:left;">
            <td style="padding:12px; font-weight:600; color:#1e293b;">{r['employee_id']}</td>
            <td style="padding:12px; color:#334155;"><b>{r['name']}</b></td>
            <td style="padding:12px; color:#64748b;">{r['department']}</td>
            <td style="padding:12px; color:#0f172a; font-family:monospace; font-size:13px;">{login_display}</td>
            <td style="padding:12px; color:#0f172a; font-family:monospace; font-size:13px;">{logout_display}</td>
            <td style="padding:12px;">
                <span style="display:inline-block; padding:4px 10px; border-radius:12px; font-size:12px; font-weight:600; {badge_style}">
                    {status}
                </span>
            </td>
        </tr>
        """

    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Daily Attendance Report - {session_name}</title>
    </head>
    <body style="margin:0; padding:0; background-color:#f8fafc; font-family:-apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; color:#334155;">
        <table width="100%" cellpadding="0" cellspacing="0" style="background-color:#f8fafc; padding:24px 0;">
            <tr>
                <td align="center">
                    <table width="680" cellpadding="0" cellspacing="0" style="background-color:#ffffff; border-radius:12px; overflow:hidden; box-shadow:0 4px 6px -1px rgba(0,0,0,0.05), 0 2px 4px -2px rgba(0,0,0,0.05); border:1px solid #e2e8f0;">
                        
                        <!-- Header Banner -->
                        <tr>
                            <td style="background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%); padding:28px 32px; text-align:left;">
                                <table width="100%" cellpadding="0" cellspacing="0">
                                    <tr>
                                        <td>
                                            <h1 style="color:#ffffff; margin:0 0 4px 0; font-size:22px; font-weight:700; letter-spacing:-0.5px;">
                                                Keya Fusion Attendance System
                                            </h1>
                                            <p style="color:#94a3b8; margin:0; font-size:14px;">
                                                Daily Attendance Report &bull; <b style="color:#38bdf8;">{session_name}</b> &bull; {formatted_date}
                                            </p>
                                        </td>
                                        <td align="right">
                                            <span style="background-color:rgba(255,255,255,0.1); color:#38bdf8; padding:6px 14px; border-radius:20px; font-size:13px; font-weight:600; border:1px solid rgba(56,189,248,0.3);">
                                                {summary['date']}
                                            </span>
                                        </td>
                                    </tr>
                                </table>
                            </td>
                        </tr>

                        <!-- KPI Stat Cards -->
                        <tr>
                            <td style="padding:24px 32px 16px 32px;">
                                <table width="100%" cellpadding="0" cellspacing="0">
                                    <tr>
                                        <td width="23%" style="background-color:#f1f5f9; border-radius:8px; padding:14px; text-align:center;">
                                            <div style="font-size:11px; color:#64748b; font-weight:600; text-transform:uppercase;">Total Staff</div>
                                            <div style="font-size:24px; color:#1e293b; font-weight:700; margin-top:4px;">{summary['total_employees']}</div>
                                        </td>
                                        <td width="2%"></td>
                                        <td width="23%" style="background-color:#ecfdf5; border-radius:8px; padding:14px; text-align:center; border:1px solid #a7f3d0;">
                                            <div style="font-size:11px; color:#065f46; font-weight:600; text-transform:uppercase;">Present</div>
                                            <div style="font-size:24px; color:#059669; font-weight:700; margin-top:4px;">{summary['present_count']}</div>
                                        </td>
                                        <td width="2%"></td>
                                        <td width="23%" style="background-color:#fef2f2; border-radius:8px; padding:14px; text-align:center; border:1px solid #fecaca;">
                                            <div style="font-size:11px; color:#991b1b; font-weight:600; text-transform:uppercase;">Absent</div>
                                            <div style="font-size:24px; color:#dc2626; font-weight:700; margin-top:4px;">{summary['absent_count']}</div>
                                        </td>
                                        <td width="2%"></td>
                                        <td width="23%" style="background-color:#eff6ff; border-radius:8px; padding:14px; text-align:center; border:1px solid #bfdbfe;">
                                            <div style="font-size:11px; color:#1e40af; font-weight:600; text-transform:uppercase;">Attendance Rate</div>
                                            <div style="font-size:24px; color:#2563eb; font-weight:700; margin-top:4px;">{summary['attendance_rate']}%</div>
                                        </td>
                                    </tr>
                                </table>
                            </td>
                        </tr>

                        <!-- Attendance Table -->
                        <tr>
                            <td style="padding:8px 32px 24px 32px;">
                                <h3 style="font-size:15px; color:#1e293b; margin:16px 0 12px 0; font-weight:600;">
                                    Attendance Log Details ({session_name})
                                </h3>
                                <div style="overflow-x:auto;">
                                    <table width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse; font-size:13px;">
                                        <thead>
                                            <tr style="background-color:#f8fafc; border-bottom:2px solid #e2e8f0; color:#475569; font-size:12px; text-transform:uppercase; text-align:left;">
                                                <th style="padding:10px 12px;">ID</th>
                                                <th style="padding:10px 12px;">Name</th>
                                                <th style="padding:10px 12px;">Dept</th>
                                                <th style="padding:10px 12px;">In Time</th>
                                                <th style="padding:10px 12px;">Out Time</th>
                                                <th style="padding:10px 12px;">Status</th>
                                            </tr>
                                        </thead>
                                        <tbody>
                                            {rows_html}
                                        </tbody>
                                    </table>
                                </div>
                            </td>
                        </tr>

                        <!-- Note about attachment -->
                        <tr>
                            <td style="padding:0 32px 24px 32px;">
                                <div style="background-color:#f8fafc; border:1px dashed #cbd5e1; border-radius:8px; padding:12px 16px; font-size:12px; color:#64748b;">
                                    📎 <b>Attachment Included:</b> Complete detailed Excel report (<code>Daily_Attendance_{summary['date']}.xlsx</code>) is attached with this email for backup and records.
                                </div>
                            </td>
                        </tr>

                        <!-- Footer -->
                        <tr>
                            <td style="background-color:#f1f5f9; padding:18px 32px; text-align:center; border-top:1px solid #e2e8f0; font-size:12px; color:#64748b;">
                                <p style="margin:0;">This is an automated attendance email notification generated by <b>Keya Fusion Face Attendance System (KEYA-FAS)</b>.</p>
                                <p style="margin:4px 0 0 0; color:#94a3b8;">Scheduled Daily Dispatches at {get_email_config()['schedule_time_display']}</p>
                            </td>
                        </tr>

                    </table>
                </td>
            </tr>
        </table>
    </body>
    </html>
    """
    return html

def send_via_brevo_api(api_key, sender_email, receiver_email, subject, html_content, excel_bytes, attachment_name):
    import urllib.request
    import json
    import base64
    
    url = "https://api.brevo.com/v3/smtp/email"
    headers = {
        "accept": "application/json",
        "api-key": api_key.strip(),
        "content-type": "application/json"
    }
    
    excel_b64 = base64.b64encode(excel_bytes).decode('utf-8')
    payload = {
        "sender": {"name": "Keya Attendance", "email": sender_email},
        "to": [{"email": receiver_email}],
        "subject": subject,
        "htmlContent": html_content,
        "attachment": [
            {
                "name": attachment_name,
                "content": excel_b64
            }
        ]
    }
    
    req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers=headers, method='POST')
    with urllib.request.urlopen(req, timeout=30) as response:
        success_msg = f"Daily attendance report successfully sent to {receiver_email} via Brevo HTTPS API."
        print(f"[EmailReporter] {success_msg}")
        return True, success_msg

def send_via_resend_api(api_key, sender_email, receiver_email, subject, html_content, excel_bytes, attachment_name):
    import urllib.request
    import json
    import base64
    
    url = "https://api.resend.com/emails"
    headers = {
        "Authorization": f"Bearer {api_key.strip()}",
        "Content-Type": "application/json"
    }
    
    excel_b64 = base64.b64encode(excel_bytes).decode('utf-8')
    from_address = os.getenv('RESEND_FROM_EMAIL', "Keya Attendance <onboarding@resend.dev>")
    payload = {
        "from": from_address,
        "to": [receiver_email],
        "subject": subject,
        "html": html_content,
        "attachments": [
            {
                "filename": attachment_name,
                "content": excel_b64
            }
        ]
    }
    
    req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers=headers, method='POST')
    with urllib.request.urlopen(req, timeout=30) as response:
        success_msg = f"Daily attendance report successfully sent to {receiver_email} via Resend HTTPS API."
        print(f"[EmailReporter] {success_msg}")
        return True, success_msg

def send_daily_attendance_email(target_date=None, to_email=None, session_name=None):
    """
    Builds the daily report and sends it to the destination email address.
    Supports Brevo / Resend HTTPS APIs (recommended for Render cloud hosting where SMTP is blocked),
    with automatic fallback to Gmail SMTP SSL/TLS.
    Returns (success: bool, message: str)
    """
    config = get_email_config()
    sender_email = config['sender']
    sender_password = config['password']
    receiver_email = to_email or config['receiver']

    if not target_date:
        target_date = datetime.date.today().strftime('%Y-%m-%d')

    if not session_name:
        current_hour = datetime.datetime.now().hour
        session_name = "Evening Summary (8:00 PM)" if current_hour >= 16 else "Morning Check-In (9:05 AM)"

    print(f"[{datetime.datetime.now()}] Preparing {session_name} email for date {target_date}...")
    
    try:
        summary, records = get_daily_attendance_data(target_date)
        html_content = generate_html_email_body(summary, records, session_name=session_name)
        excel_bytes = generate_daily_excel_bytes(summary, records)
    except Exception as e:
        err_msg = f"Failed to gather attendance data: {str(e)}"
        print(f"[EmailReporter] Error: {err_msg}")
        return False, err_msg

    attachment_name = f"Daily_Attendance_{target_date}.xlsx"
    subject = f"Keya Fusion Attendance [{session_name}] - {target_date} ({summary['present_count']}/{summary['total_employees']} Present)"

    # 1. Check for HTTPS API keys (Brevo / Resend) - Essential for Render Cloud Hosting
    brevo_api_key = os.getenv('BREVO_API_KEY')
    if brevo_api_key:
        try:
            return send_via_brevo_api(brevo_api_key, sender_email, receiver_email, subject, html_content, excel_bytes, attachment_name)
        except Exception as e_brevo:
            print(f"[EmailReporter] Brevo HTTPS API error: {e_brevo}")

    resend_api_key = os.getenv('RESEND_API_KEY')
    if resend_api_key:
        try:
            return send_via_resend_api(resend_api_key, sender_email, receiver_email, subject, html_content, excel_bytes, attachment_name)
        except Exception as e_resend:
            print(f"[EmailReporter] Resend HTTPS API error: {e_resend}")

    # 2. Build MIME message for direct SMTP
    msg = MIMEMultipart('mixed')
    msg['Subject'] = subject
    msg['From'] = f"Keya Attendance System <{sender_email}>"
    msg['To'] = receiver_email

    # Add HTML body
    part_html = MIMEText(html_content, 'html', 'utf-8')
    msg.attach(part_html)

    # Add Excel Attachment
    part_attachment = MIMEBase('application', 'vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    part_attachment.set_payload(excel_bytes)
    encoders.encode_base64(part_attachment)
    part_attachment.add_header('Content-Disposition', f'attachment; filename="{attachment_name}"')
    msg.attach(part_attachment)

    # Attempt sending via SMTP SSL (port 465) or TLS (port 587)
    smtp_server = 'smtp.gmail.com'
    error_log = []

    # Try SSL port 465
    try:
        with smtplib.SMTP_SSL(smtp_server, 465, timeout=20) as server:
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, [receiver_email], msg.as_string())
        success_msg = f"Daily attendance report successfully sent to {receiver_email} via SSL:465."
        print(f"[EmailReporter] {success_msg}")
        return True, success_msg
    except Exception as e1:
        error_log.append(f"SSL:465 attempt failed: {str(e1)}")

    # Try TLS port 587
    try:
        with smtplib.SMTP(smtp_server, 587, timeout=20) as server:
            server.starttls()
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, [receiver_email], msg.as_string())
        success_msg = f"Daily attendance report successfully sent to {receiver_email} via TLS:587."
        print(f"[EmailReporter] {success_msg}")
        return True, success_msg
    except Exception as e2:
        error_log.append(f"TLS:587 attempt failed: {str(e2)}")

    full_error = " | ".join(error_log)
    if "Network is unreachable" in full_error or "101" in full_error:
        hint = (
            "Render Cloud Hosting blocks direct SMTP ports (25, 465, 587) to prevent spam. "
            "To send emails from Render, get a free API Key from Brevo (https://brevo.com) or Resend (https://resend.com) "
            "and add BREVO_API_KEY or RESEND_API_KEY in Render Dashboard -> Environment."
        )
        final_msg = f"{full_error}. [Render Notice: {hint}]"
    elif "535" in full_error or "Username and Password not accepted" in full_error:
        hint = (
            "Gmail SMTP authentication failed. For security, Gmail requires a 16-character 'Google App Password'. "
            "Please go to Google Account -> Security -> 2-Step Verification -> App Passwords, "
            "generate an App Password and put it in .env / Render Environment as REPORT_SENDER_PASSWORD."
        )
        final_msg = f"{full_error}. TIP: {hint}"
    else:
        final_msg = f"Failed to send email: {full_error}"

    print(f"[EmailReporter] Error: {final_msg}")
    return False, final_msg
